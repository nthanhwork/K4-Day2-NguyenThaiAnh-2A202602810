"""Build the submission from saved predictions and immutable experiment evidence.

No model construction, training, temperature fitting or test inference occurs here.
Use --snapshot once on the original workspace, then rerun from evidence on any CPU.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[3]
SUB = Path(__file__).resolve().parents[1]
EVIDENCE = SUB / "evidence"
FINAL = Path("runs/finalization/ConvNeXt_Final_v1")
SESSION = Path("runs/inference/I_T04_seed0")
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".cache/matplotlib"))
sys.path.insert(0, str(ROOT))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from PIL import Image

from eval import check_against_csv, compute_metrics, load_names, read_pred


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text())


def write_json(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n")


def text_style(text):
    """Space prose labels and units without changing experiment IDs or file paths."""
    text = re.sub(r"\b(F01|T00|T04|B03)seed", r"\1 seed", text)
    prefixes = "train|val|test|seed|fold|batch|input|resize|crop|warmup|epoch|recipe|support|clip|Table|Negative|CE|WD|Python|torchvision|torch|timm|numpy|CUDA|initialCE|evalCE|ln"
    text = re.sub(r"(?<![\w/])(" + prefixes + r")(?=\d)", r"\1 ", text)
    text = re.sub(r"(?<=\d)(?=(?:ms|GB|MB|epoch|seed|sheet|steps|updates|lớp|ảnh|warmup)\b)", " ", text)
    return re.sub(r"(?<=[,;])(?=[A-Za-zÀ-ỿ])", " ", text)


def snapshot():
    """Copy small original logs verbatim; leave checkpoints and live protocols alone."""
    completion = read_json(ROOT / FINAL / "completion.json")
    if not completion["section10_complete"]:
        raise ValueError("Final evaluation must be complete before report export")
    original = []
    for group in ("B01", "B02", "B03", "B04", "B05", "T00", "T01", "T02", "T03",
                  "T04", "T05", "T06", "T07", "T08", "F01", "ablation_convnext_tiny",
                  "inference/I_T04_seed0", "finalization/ConvNeXt_Final_v1"):
        folder = ROOT / "runs" / group
        for path in sorted(folder.rglob("*")):
            if path.is_file() and path.suffix in {".json", ".csv", ".log", ".md", ".sha256"}:
                original.append(path)
    original += list((ROOT / "data/labels").glob("*.csv"))
    records = []
    for source in original:
        relative = source.relative_to(ROOT)
        target_relative = Path("labels") / source.name if relative.parts[0] == "data" else relative
        target = EVIDENCE / target_relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, target)
        records.append({"source": str(relative), "artifact": str(target.relative_to(SUB)),
                        "sha256": sha(target), "bytes": target.stat().st_size})
    write_json(EVIDENCE / "source_index.json", {
        "schema": 1, "original_workspace": str(ROOT), "verbatim_copies": True,
        "checkpoints_included": False, "raw_arrays_included": False, "files": records,
    })
    (EVIDENCE / "README.md").write_text(
        "# Bằng chứng thực nghiệm\n\n"
        "Các file trong `runs/` và `labels/` là bản sao nguyên byte của log, cấu hình, "
        "protocol, evaluator và CSV nguồn. `source_index.json` ghi SHA256 và đường dẫn nguồn. "
        "Đường dẫn tuyệt đối trong protocol phản ánh máy thực nghiệm; giữ nguyên để bảo toàn checksum.\n\n"
        "Các bản sao phục vụ kiểm tra và tạo báo cáo, không đặt trở lại `runs/` để resume "
        "trên máy khác. Checkpoint và raw NumPy cache được giữ ở máy gốc, không commit. "
        "Dự đoán CSV đầy đủ nằm tại `../predictions/`. Máy mới có thể kiểm tra metric "
        "và tạo lại Excel/báo cáo bằng CPU; tái huấn luyện dùng một checkout sạch.\n",
        encoding="utf-8")


def verify_snapshot():
    index = read_json(EVIDENCE / "source_index.json")
    for item in index["files"]:
        if sha(SUB / item["artifact"]) != item["sha256"]:
            raise ValueError(f"Evidence changed: {item['artifact']}")
    source_hashes = read_json(SUB / "code/data_sources.json")["files"]
    for name, info in source_hashes.items():
        if sha(EVIDENCE / "labels" / name) != info["sha256"]:
            raise ValueError(f"Original CSV changed: {name}")
    if sha(ROOT / "eval.py") != read_json(EVIDENCE / FINAL / "completion.json")["evaluator_sha256"]:
        raise ValueError("Official evaluator differs from the one used for final evaluation")


def prediction(name, split):
    path = SUB / "predictions" / name
    pred = read_pred(str(path))
    reference = EVIDENCE / "labels" / f"{split}_subset0.csv"
    check_against_csv(pred, str(reference), split)
    ref = pd.read_csv(reference)
    if pred.filenames.tolist() != ref.Filename.tolist():
        raise ValueError(f"Prediction row order differs: {name}")
    return pred, compute_metrics(pred.y_true, pred.y_pred, pred.probs)


def metric_row(metrics):
    return {key: float(metrics[key]) for key in ("top1", "macro_f1", "balanced_acc", "ece", "nll")}


def latency_rows():
    measured = read_json(SUB / "code/validation/backbone_latency.json")["measurements"]
    raw = [(item["exp_id"], "backbone", item) for item in measured]
    for method in ("I00", "I01", "I02", "I07", "I08"):
        raw.extend((method, "inference", item) for item in read_json(EVIDENCE / SESSION / method / "latency.json"))
    rows = []
    for exp_id, stage, item in raw:
        values = np.asarray(item["samples_ms"], dtype=float)
        if len(values) != item["n"] or item["warmup"] < 10 or item["n"] < 50 or not np.isfinite(values).all() or (values <= 0).any():
            raise ValueError(f"Invalid timing samples: {exp_id}")
        for q in (50, 95, 99):
            if not np.isclose(np.percentile(values, q), item[f"p{q}"]):
                raise ValueError(f"Invalid percentile: {exp_id}")
        rows.append({"exp_id": exp_id, "stage": stage, "measurement_id": exp_id,
                     **{key: item[key] for key in ("gpu", "dtype", "batch", "bn_fused", "k_views", "warmup", "n", "scope")},
                     **{f"p{q}_ms": item[f"p{q}"] for q in (50, 95, 99)},
                     "images_per_s": item["images_per_s"], "representative_for_3_seeds": False})
    for role, source in (("T00", "B03"), ("F01", "I07")):
        for row in list(rows):
            if row["exp_id"] == source:
                rows.append({**row, "exp_id": role, "stage": "final representative",
                             "representative_for_3_seeds": True})
    return pd.DataFrame(rows)


def tables():
    latency = latency_rows()
    def measured(exp_id, batch=1):
        return latency[(latency.exp_id == exp_id) & (latency.batch == batch)].iloc[0]
    backbone_rows = []
    for exp_id in ("B01", "B02", "B03", "B04", "B05"):
        folder = EVIDENCE / "runs" / exp_id / "seed0"
        cfg, model, summary = (read_json(folder / name) for name in ("config.json", "model.json", "summary.json"))
        _, metrics = prediction(f"{exp_id}_seed0_val.csv", "val")
        for key in ("top1", "macro_f1", "ece"):
            np.testing.assert_allclose(metrics[key], summary["val"][key], atol=2e-7)
        timing = measured(exp_id)
        backbone_rows.append({"exp_id": exp_id, "backbone": cfg["backbone"],
            "pretrained_tag": model["pretrained_cfg"]["tag"], "params_M": summary["params_m"],
            "GMAC": summary["gmacs"], "img_size_px": cfg["img_size"], "epochs": cfg["epochs"],
            "seed": cfg["seed"], "n_seeds": 1, "best_epoch": summary["best_epoch"],
            "macro_f1_val": metrics["macro_f1"], "top1_val": metrics["top1"],
            "train_s_per_epoch_observed": summary["train_seconds_per_epoch"],
            "latency_p50_batch1_ms": timing.p50_ms, "latency_p95_batch1_ms": timing.p95_ms,
            "throughput_batch32_img_s": measured(exp_id, 32).images_per_s,
            "curve": f"curves/{exp_id}_seed0_{cfg['backbone']}.png",
            "evidence": str(folder.relative_to(SUB)),
            "notes": "Training time may include GPU contention; GMAC excludes fvcore unsupported operators; different pretraining recipes"})
    backbones = pd.DataFrame(backbone_rows)
    training = pd.read_csv(EVIDENCE / "runs/ablation_convnext_tiny/training_ablation.csv")
    training = training.drop(columns=["config_json", "run_path"], errors="ignore")
    for index, row in training.iterrows():
        source_id = "B03" if row.exp_id == "T00" else row.exp_id
        _, metrics = prediction(f"{source_id}_seed0_val.csv", "val")
        training.loc[index, "macro_f1_val"] = metrics["macro_f1"]
        training.loc[index, "top1_val"] = metrics["top1"]
        training.loc[index, "chinee_apple_f1_val"] = metrics["f1"][0]
        training.loc[index, "snake_weed_f1_val"] = metrics["f1"][7]
        training.loc[index, "curve"] = f"curves/{row.exp_id}_seed0_convnext_tiny.png"
    training["n_seeds"] = 1
    inference = pd.read_csv(EVIDENCE / SESSION / "inference_results.csv")
    inference = inference.rename(columns={"method": "exp_id"}).drop(columns=["prediction_path"], errors="ignore")
    inference["temperature"] = [read_json(EVIDENCE / SESSION / "temperature.json")["T"] if m == "I07" else 1.0 for m in inference.exp_id]
    inference["n_seeds"] = 1
    inference["curve"] = "curves/inference/I_T04_seed0/f1_vs_latency_p95_detail.png"
    inference["latency_scope"] = measured("I00").scope
    final_rows, class_rows = [], []
    names = load_names(str(EVIDENCE / "labels/labels.csv"))
    for role in ("T00", "F01"):
        all_metrics = []
        for seed in (0, 1, 2):
            folder = EVIDENCE / FINAL / role / f"seed{seed}"
            prepared = read_json(folder / "prepared.json")
            _, val = prediction(f"{role}_seed{seed}_val.csv", "val")
            _, test = prediction(f"{role}_seed{seed}_test.csv", "test")
            row = {"exp_id": role, "configuration": "ConvNeXt-Tiny + " + ("CutMix + I07" if role == "F01" else "basic CE + I00"),
                "seed": seed, "n_seeds": 1, "source_id": prepared["source_id"], "seed0_reuse": prepared["seed0_reuse"],
                "temperature_fit_val": prepared["temperature"], "best_epoch": prepared["best_epoch"],
                "macro_f1_val": val["macro_f1"], "macro_f1_val_std": None,
                **{f"{key}_test": float(value) for key, value in metric_row(test).items()},
                **{f"{key}_test_std": None for key in metric_row(test)}, "n_test_per_seed": len(read_pred(str(SUB / "predictions" / f"{role}_seed{seed}_test.csv")).y_true),
                "latency_p95_batch1_ms_representative": measured(role).p95_ms,
                "curve": f"curves/{role}_seed{seed}_convnext_tiny.png"}
            final_rows.append(row)
            all_metrics.append((val, test))
            for label, name in enumerate(names):
                class_rows.append({"exp_id": role, "seed": seed, "n_seeds": 1, "class_id": label,
                    "class": name, "support_per_seed": int(test["support"][label]),
                    **{key: float(test[key][label]) for key in ("precision", "recall", "f1")},
                    **{key + "_std": None for key in ("precision", "recall", "f1")}})
        official = read_json(EVIDENCE / FINAL / "eval_out" / f"{role}_summary.json")
        aggregate = {**final_rows[-1], "seed": "mean", "n_seeds": 3, "source_id": "see per-seed rows",
                     "seed0_reuse": None, "temperature_fit_val": None, "best_epoch": None, "curve": "see per-seed rows",
                     "macro_f1_val": float(np.mean([m[0]["macro_f1"] for m in all_metrics])),
                     "macro_f1_val_std": float(np.std([m[0]["macro_f1"] for m in all_metrics], ddof=1))}
        for key in metric_row(all_metrics[0][1]):
            values = [m[1][key] for m in all_metrics]
            mean, std = float(np.mean(values)), float(np.std(values, ddof=1))
            np.testing.assert_allclose([mean, std], [official[key]["mean"], official[key]["std"]], atol=1e-12)
            aggregate[f"{key}_test"] = mean
            aggregate[f"{key}_test_std"] = std
        final_rows.append(aggregate)
        for label, name in enumerate(names):
            class_rows.append({"exp_id": role, "seed": "mean", "n_seeds": 3, "class_id": label,
                "class": name, "support_per_seed": int(all_metrics[0][1]["support"][label]),
                **{key: official[key]["mean"][label] for key in ("precision", "recall", "f1")},
                **{key + "_std": official[key]["std"][label] for key in ("precision", "recall", "f1")}})
    final = pd.DataFrame(final_rows)
    summary_rows = []
    for row in backbones.to_dict("records"):
        summary_rows.append({"exp_id": row["exp_id"], "stage": "backbone", "macro_f1_val": row["macro_f1_val"],
            "top1_val": row["top1_val"], "n_seeds": 1, "p95_batch1_ms": row["latency_p95_batch1_ms"],
            "GMAC": row["GMAC"], "params_M": row["params_M"], "notes": "Single-seed screening"})
    for row in training.to_dict("records"):
        if row["exp_id"] == "T00":
            continue  # Same measurement as B03, already included.
        summary_rows.append({"exp_id": row["exp_id"], "stage": "training", "macro_f1_val": row["macro_f1_val"],
            "top1_val": row["top1_val"], "n_seeds": 1, "p95_batch1_ms": None,
            "GMAC": backbones.loc[backbones.exp_id == "B03", "GMAC"].iloc[0],
            "params_M": backbones.loc[backbones.exp_id == "B03", "params_M"].iloc[0],
            "notes": "Recipe latency not separately measured; use Inference sheet for measured T04"})
    for row in inference.to_dict("records"):
        summary_rows.append({"exp_id": row["exp_id"], "stage": "inference on T04", "macro_f1_val": row["macro_f1"],
            "top1_val": row["top1"], "n_seeds": 1, "p95_batch1_ms": row["latency_p95_ms"],
            "GMAC": None, "params_M": backbones.loc[backbones.exp_id == "B03", "params_M"].iloc[0],
            "notes": "Same checkpoint across methods; I07 selected on val, I08 test not measured"})
    for row in final[final.seed == "mean"].to_dict("records"):
        summary_rows.append({"exp_id": row["exp_id"], "stage": "final mean", "macro_f1_val": row["macro_f1_val"],
            "top1_val": float(np.mean([prediction(f"{row['exp_id']}_seed{s}_val.csv", "val")[1]["top1"] for s in (0, 1, 2)])),
            "n_seeds": 3, "p95_batch1_ms": row["latency_p95_batch1_ms_representative"],
            "GMAC": backbones.loc[backbones.exp_id == "B03", "GMAC"].iloc[0],
            "params_M": backbones.loc[backbones.exp_id == "B03", "params_M"].iloc[0],
            "notes": "Mean val over fixed seeds; latency representative seed0, models are independent"})
    summary = pd.DataFrame(summary_rows).sort_values(["macro_f1_val", "exp_id"], ascending=[False, True]).head(10).reset_index(drop=True)
    summary.insert(0, "rank", range(1, len(summary) + 1))
    return {"Backbones": backbones, "Training": training, "Inference": inference,
            "Final": final, "PerClass": pd.DataFrame(class_rows), "Latency": latency, "Summary": summary}


def workbook(frames):
    output = SUB / "results.xlsx"
    notes = {
        "Backbones": "Fold0, seed0. All metrics are fractions. Train times are observed and may include shared GPU. GMAC uses fvcore, unsupported ops excluded.",
        "Training": "T00 aliases B03; T01-T07 each change one factor. T08 combines T04 and T07. Delta_pp is percentage points. Screening has one seed, std unavailable.",
        "Inference": "Validation only, T04/seed0. Latency includes GPU crop/views, model, aggregation/calibration and softmax; excludes CPU/disk/H2D. Throughput uses batch32.",
        "Final": "Seeds0,1,2; sample std ddof=1. Blank std means unavailable for one seed. Test3507/seed. F01seed0 aliases T04; T00seed0 aliases B03. Latency is representative, not remeasured across seeds.",
        "PerClass": "Test support is per seed, not summed over 3 seeds. Mean/std are computed across seed metrics, not from pooled confusion.",
        "Latency": "RTX3060,10 warmup,100 synchronized timings per measurement. Images/s=batch divided by p50 seconds. Final rows alias named measurements; no independent timing claim.",
        "Summary": "Top10 by validation macro-F1 only. Some entries share checkpoints. Blank latency means not measured for that recipe; final val is a 3-seed mean. See Final for test results.",
    }
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, frame in frames.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False
            ws.row_dimensions[1].height = 34
            ws["A1"].comment = Comment(notes[name], "Nguyen Thai Anh")
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="18324F")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for column in ws.columns:
                title = str(column[0].value)
                ws.column_dimensions[column[0].column_letter].width = min(40, max(13, len(title) + 2))
                percent = any(t in title for t in ("macro_f1", "top1", "balanced_acc", "precision", "recall", "f1")) and "pp" not in title
                for cell in column[1:]:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if isinstance(cell.value, float):
                        cell.number_format = "0.00%" if percent else "0.0000"
                    if title in {"evidence", "curve"} and cell.value and str(cell.value).startswith(("curves/", "evidence/")):
                        cell.hyperlink = str(cell.value)
                        cell.font = Font(color="0563C1", underline="single")
            if name == "Backbones":
                best = frame.macro_f1_val.idxmax() + 2
            elif name == "Training":
                best = frame.macro_f1_val.idxmax() + 2
            elif name == "Inference":
                best = frame.index[frame.exp_id == "I07"][0] + 2
            elif name == "Final":
                best = frame.index[(frame.exp_id == "F01") & (frame.seed == "mean")][0] + 2
            elif name == "Summary":
                best = 2
            else:
                best = None
            if best:
                for cell in ws[best]:
                    cell.fill = PatternFill("solid", fgColor="DDF1E4")
                    cell.font = Font(bold=True)
            ws.print_options.horizontalCentered = True
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.print_title_rows = "1:1"
    for name, frame in frames.items():
        (EVIDENCE / "tables").mkdir(parents=True, exist_ok=True)
        frame.to_csv(EVIDENCE / "tables" / f"{name}.csv", index=False)
    return output


def save_figure(fig, name):
    output = SUB / "curves/final_analysis" / name
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return output


def reliability(pred):
    conf = pred.probs.max(1)
    correct = (pred.y_pred == pred.y_true).astype(float)
    ids = np.clip(np.ceil(conf * 15).astype(int) - 1, 0, 14)
    return pd.DataFrame([{"bin": i, "lower": i / 15, "upper": (i + 1) / 15,
        "count": int((ids == i).sum()),
        "confidence": float(conf[ids == i].mean()) if (ids == i).any() else np.nan,
        "accuracy": float(correct[ids == i].mean()) if (ids == i).any() else np.nan} for i in range(15)])


def figures(frames, images_dir):
    names = load_names(str(EVIDENCE / "labels/labels.csv"))
    b = frames["Backbones"]
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2), sharey=True)
    for row in b.itertuples():
        for ax, x in zip(axes, (row.params_M, row.latency_p95_batch1_ms)):
            ax.scatter(x, row.macro_f1_val * 100)
            ax.annotate(row.exp_id, (x, row.macro_f1_val * 100), xytext=(4, 5), textcoords="offset points")
    axes[0].set_xlabel("Parameters (million)")
    axes[1].set_xlabel("Isolated GPU pipeline p95, batch1 (ms)")
    axes[0].set_ylabel("Validation macro-F1 (%)")
    fig.suptitle("B01-B05 | fold0, seed0 | different pretrained tags")
    fig.tight_layout()
    save_figure(fig, "B_backbones_tradeoff.png")
    for role, source in (("T00", "B03"), ("F01", "T04")):
        history = pd.read_csv(EVIDENCE / "runs" / source / "seed0/history.csv")
        fig, axes = plt.subplots(1, 3, figsize=(12, 3.5))
        axes[0].plot(history.epoch, history.train_loss, label="train")
        axes[0].plot(history.epoch, history.val_loss, label="val")
        axes[0].set_ylabel("Loss (training objective / val CE)")
        axes[0].legend()
        axes[1].plot(history.epoch, history.val_macro_f1)
        axes[1].set_ylabel("Validation macro-F1")
        axes[2].plot(history.epoch, history.lr)
        axes[2].set_ylabel("Backbone LR at epoch end")
        for ax in axes:
            ax.set_xlabel("Epoch")
        fig.suptitle(f"{role} seed0 ConvNeXt-Tiny | reuses {source} seed0; no new training")
        fig.tight_layout()
        fig.savefig(SUB / "curves" / f"{role}_seed0_convnext_tiny.png", dpi=150, bbox_inches="tight")
        plt.close(fig)
    for role in ("T00", "F01"):
        matrices = [compute_metrics(p.y_true, p.y_pred, p.probs)["confusion"] for p in
                    [prediction(f"{role}_seed{s}_test.csv", "test")[0] for s in (0, 1, 2)]]
        cm = np.sum(matrices, axis=0)
        # A pooled count is explicit; each image appears once per seed.
        if int(cm.sum()) != 3 * 3507:
            raise ValueError("Invalid pooled test support")
        normalized = cm / cm.sum(1, keepdims=True) * 100
        fig, axes = plt.subplots(1, 2, figsize=(13, 5.5))
        for ax, values, title in zip(axes, (cm, normalized), ("Counts summed over 3 seeds (N=10521)", "Row-normalized (%)")):
            ax.imshow(values, cmap="Blues")
            ax.set_xticks(range(9), names, rotation=55, ha="right", fontsize=8)
            ax.set_yticks(range(9), names, fontsize=8)
            ax.set_xlabel("Predicted")
            ax.set_ylabel("True")
            ax.set_title(title)
            for i in range(9):
                for j in range(9):
                    label = str(int(values[i, j])) if ax is axes[0] else f"{values[i, j]:.1f}"
                    ax.text(j, i, label, ha="center", va="center", fontsize=7,
                            color="white" if values[i, j] > values.max() / 2 else "black")
        fig.suptitle(f"{role}: fixed test fold0, seeds0/1/2 (not an ensemble)")
        fig.tight_layout()
        save_figure(fig, f"{role}_confusion_test.png")
    per_class = frames["PerClass"]
    fig, ax = plt.subplots(figsize=(10, 4))
    x = np.arange(9)
    for role, offset in (("T00", -.18), ("F01", .18)):
        part = per_class[(per_class.exp_id == role) & (per_class.seed == "mean")].sort_values("class_id")
        ax.bar(x + offset, part.f1 * 100, width=.36, yerr=part.f1_std * 100, label=role, capsize=3)
    ax.set_xticks(x, names, rotation=35, ha="right")
    ax.set_ylim(85, 100)
    ax.set_ylabel("Test F1 (%) | mean +/- sample std, 3 seeds")
    ax.legend()
    fig.tight_layout()
    save_figure(fig, "F01_vs_T00_per_class_test.png")
    fig, axes = plt.subplots(1, 3, figsize=(12, 3.8))
    ece_rows = []
    for seed, ax in enumerate(axes):
        ax.plot([0, 1], [0, 1], "--", color="gray")
        for tag, label in (("F01uncal", "Before TS"), ("F01", "After TS")):
            pred, metric = prediction(f"{tag}_seed{seed}_test.csv", "test")
            bins = reliability(pred)
            bins.to_csv(EVIDENCE / "tables" / f"{tag}_seed{seed}_test_reliability.csv", index=False)
            nonempty = bins[bins["count"] > 0]
            ax.plot(nonempty.confidence, nonempty.accuracy, "o-", label=f"{label}: ECE={metric['ece']:.4f}", markersize=3)
            ece_rows.append({"tag": tag, "seed": seed, "ece": metric["ece"], "nll": metric["nll"]})
        ax.set_title(f"Seed {seed} | 15 bins, T fitted on val")
        ax.set_xlabel("Mean confidence in nonempty bin")
        ax.set_ylabel("Accuracy")
        ax.legend(fontsize=8)
        ax.set_xlim(0, 1)
        ax.set_ylim(0, 1.03)
    fig.tight_layout()
    save_figure(fig, "F01_test_reliability_before_after.png")
    pd.DataFrame(ece_rows).to_csv(EVIDENCE / "tables/test_calibration.csv", index=False)
    pred, _ = prediction("F01_seed0_test.csv", "test")
    errors = pd.DataFrame({"Filename": pred.filenames, "y_true": pred.y_true, "y_pred": pred.y_pred,
                           "confidence": pred.probs.max(1)})
    errors = errors[errors.y_true != errors.y_pred].sort_values(["confidence", "Filename"], ascending=[False, True])
    errors["true_class"] = errors.y_true.map(dict(enumerate(names)))
    errors["predicted_class"] = errors.y_pred.map(dict(enumerate(names)))
    errors.to_csv(EVIDENCE / "tables/F01_seed0_test_errors.csv", index=False)
    pairs = errors.groupby(["true_class", "predicted_class"]).size().reset_index(name="n").sort_values("n", ascending=False)
    pairs.to_csv(EVIDENCE / "tables/F01_seed0_error_pairs.csv", index=False)
    hard_pair = errors[errors.y_true.isin([0, 7]) & errors.y_pred.isin([0, 7])]
    hard_other = errors[errors.y_true.isin([0, 7])]
    chosen = pd.concat([hard_pair[hard_pair.y_true == 0].head(2), hard_pair[hard_pair.y_true == 7].head(2),
                        hard_other.head(6), errors.head(12)]).drop_duplicates("Filename").head(12)
    chosen.to_csv(EVIDENCE / "tables/F01_seed0_error_gallery.csv", index=False)
    paths = [Path(images_dir) / name for name in chosen.Filename]
    if all(path.exists() for path in paths):
        fig, axes = plt.subplots(3, 4, figsize=(12, 11.5), layout="constrained")
        for ax, (_, row) in zip(axes.flat, chosen.iterrows()):
            with Image.open(Path(images_dir) / row.Filename) as image:
                ax.imshow(image.convert("RGB"))
            ax.set_title(f"{row.Filename}\nTrue: {row.true_class}\nPred: {row.predicted_class} ({row.confidence:.3f})", fontsize=8)
            ax.axis("off")
        for ax in list(axes.flat)[len(chosen):]:
            ax.axis("off")
        fig.suptitle("F01 seed0 test errors | hard pairs first, then highest confidence; no new inference")
        save_figure(fig, "F01_seed0_error_gallery.png")
    elif not (SUB / "curves/final_analysis/F01_seed0_error_gallery.png").exists():
        raise FileNotFoundError("Original images needed once to build error gallery; use --images-dir")


def markdown_table(frame, columns, percent=()):
    lines = ["| " + " | ".join(columns) + " |", "| " + " | ".join("---" for _ in columns) + " |"]
    for row in frame.to_dict("records"):
        values = []
        for col in columns:
            value = row[col]
            if pd.isna(value):
                text = "—"
            elif col in percent:
                text = f"{value * 100:.4f}"
            elif isinstance(value, (float, np.floating)):
                text = f"{value:.4f}"
            else:
                text = str(value)
            values.append(text.replace("|", "/").replace("\n", " "))
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def export(images_dir=ROOT / "data", make_snapshot=False):
    if make_snapshot:
        snapshot()
    verify_snapshot()
    frames = tables()
    workbook(frames)
    figures(frames, images_dir)
    from submission_report import write_report
    write_report(frames)
    from submission_pdf import write_pdf
    write_pdf(frames, Path(images_dir))
    return frames


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot", action="store_true", help="copy original completed run evidence first")
    parser.add_argument("--images-dir", type=Path, default=ROOT / "data")
    args = parser.parse_args()
    frames = export(args.images_dir, args.snapshot)
    print("Export complete:", {name: len(frame) for name, frame in frames.items()})
